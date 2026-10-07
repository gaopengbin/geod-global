"""Regenerate owned Office fixtures. Libraries are authoring tools, not app dependencies."""
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET
from docx import Document
from openpyxl import Workbook

out = Path(__file__).resolve().parent
document = Document()
document.core_properties.author = "GeoD"
document.add_heading("GeoD 文档检查", 0)
document.add_paragraph("北京 🌍 & <GeoD>")
document.add_paragraph("<script>window.officeInjected=true</script>")
table = document.add_table(rows=2, cols=2)
for cells, values in zip(table.rows, [("城市", "经度"), ("北京", "116.4")]):
    for cell, value in zip(cells.cells, values):
        cell.text = value
document.sections[0].header.paragraphs[0].text = "GeoD header"
document.save(out / "notes.docx")

workbook = Workbook()
workbook.properties.creator = "GeoD"
workbook.active.title = "数据"
workbook.active.append(["城市", "值"])
workbook.active.append(["北京", 1.25])
workbook.active.append(["上海", 2.5])
workbook.active["B4"] = "=SUM(B2:B3)"
workbook.create_sheet("说明").append(["公式不会执行 & 原文件保持不变"])
buffer = BytesIO()
workbook.save(buffer)
with ZipFile(buffer) as source:
    parts = {name: source.read(name) for name in source.namelist()}
# Exercise shared strings and inline strings in one ordinary workbook.
sheet_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
root = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
cell = root.find(f".//{{{sheet_ns}}}c[@r='A1']")
for child in list(cell):
    cell.remove(child)
cell.set("t", "s")
ET.SubElement(cell, f"{{{sheet_ns}}}v").text = "0"
parts["xl/worksheets/sheet1.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
parts["xl/sharedStrings.xml"] = f'<sst xmlns="{sheet_ns}" count="1" uniqueCount="1"><si><t>城市</t><rPh sb="0" eb="2"><t>chengshi</t></rPh></si></sst>'.encode()
rels = ET.fromstring(parts["xl/_rels/workbook.xml.rels"])
ET.SubElement(rels, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship", {
    "Id": "rIdShared", "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings", "Target": "sharedStrings.xml"})
parts["xl/_rels/workbook.xml.rels"] = ET.tostring(rels, encoding="utf-8", xml_declaration=True)
types = ET.fromstring(parts["[Content_Types].xml"])
ET.SubElement(types, "{http://schemas.openxmlformats.org/package/2006/content-types}Override", {
    "PartName": "/xl/sharedStrings.xml", "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"})
parts["[Content_Types].xml"] = ET.tostring(types, encoding="utf-8", xml_declaration=True)
with ZipFile(out / "data.xlsx", "w", ZIP_DEFLATED) as target:
    for name, data in parts.items():
        target.writestr(name, data)

relationship_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
office_rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
parts = {
    "[Content_Types].xml": '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/><Override PartName="/ppt/slides/slide1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/><Override PartName="/ppt/slides/slide2.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/></Types>',
    "_rels/.rels": f'<Relationships xmlns="{relationship_ns}"><Relationship Id="rId1" Type="{office_rel}/officeDocument" Target="ppt/presentation.xml"/></Relationships>',
    "ppt/presentation.xml": f'<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:r="{office_rel}"><p:sldIdLst><p:sldId r:id="rId2" id="257"/><p:sldId id="256" r:id="rId1"/></p:sldIdLst><p:sldSz cx="9144000" cy="6858000"/><p:notesSz cx="6858000" cy="9144000"/></p:presentation>',
    "ppt/_rels/presentation.xml.rels": f'<Relationships xmlns="{relationship_ns}"><Relationship Id="rId1" Type="{office_rel}/slide" Target="slides/slide1.xml"/><Relationship Id="rId2" Type="{office_rel}/slide" Target="slides/slide2.xml"/></Relationships>',
}
for number, text in [(1, "Second in presentation order"), (2, "第一张 &amp; GeoD")]:
    parts[f"ppt/slides/slide{number}.xml"] = '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:cSld><p:spTree><p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr/><p:sp><p:nvSpPr><p:cNvPr id="2" name="Title"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>' + text + '</a:t></a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld><p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr></p:sld>'
with ZipFile(out / "slides.pptx", "w", ZIP_DEFLATED) as target:
    for name, data in parts.items():
        target.writestr(name, data.encode("utf-8"))
print("Wrote three owned OOXML fixtures")
